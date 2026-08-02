"""Export cleaned/enhanced datasets in common portable formats."""

from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

DATA_EXPORT_FORMATS = {"csv", "xlsx", "json"}


def _sanitize_excel_value(value: Any) -> Any:
    """Prevent spreadsheet formula injection for user-controlled values."""
    if isinstance(value, str) and value and value[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{value}"
    return value


def _json_safe(value: Any) -> Any:
    if pd.isna(value):
        return None
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:  # noqa: BLE001
            pass
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    return value


def export_dataframe(frame: pd.DataFrame, export_format: str) -> tuple[bytes, str, str]:
    """Return bytes, media type, and file extension for a cleaned dataset."""
    if export_format not in DATA_EXPORT_FORMATS:
        raise ValueError(f"Unsupported data format: {export_format}")

    if export_format == "csv":
        return (
            frame.to_csv(index=False).encode("utf-8"),
            "text/csv",
            "csv",
        )

    if export_format == "json":
        records = [
            {str(key): _json_safe(value) for key, value in row.items()}
            for row in frame.to_dict(orient="records")
        ]
        return (
            json.dumps({"data": records}, indent=2, ensure_ascii=False).encode("utf-8"),
            "application/json",
            "json",
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Enhanced Data"
    sheet["A1"] = "SOFIDR enhanced dataset"
    sheet["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="0B3D2E")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(frame.shape[1], 1))

    for row_index, row in enumerate(
        dataframe_to_rows(frame, index=False, header=True),
        start=2,
    ):
        for column_index, value in enumerate(row, start=1):
            sheet.cell(row=row_index, column=column_index, value=_sanitize_excel_value(value))

    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B7F5A")
    sheet.freeze_panes = "A3"
    for index, column in enumerate(frame.columns, start=1):
        letter = sheet.cell(row=2, column=index).column_letter
        sheet.column_dimensions[letter].width = min(max(len(str(column)) + 4, 12), 40)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return (
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    )


def safe_filename_stem(name: str | None, fallback: str = "dataset") -> str:
    stem = Path(name or fallback).stem
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("._")
    return cleaned or fallback
